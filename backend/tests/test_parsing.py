"""Resume parsing, document extraction/validation and chunking."""

import io
import zipfile

import pytest

from app.database.models import CandidateProfile, Project, Skill
from app.rag.chunking import profile_chunks, sentence_windows
from app.services.documents import DocumentError, extract_document
from app.services.resume_parser import parse_resume, split_sections
from tests.conftest import SAMPLE_RESUME

TEXT = SAMPLE_RESUME.decode()


def _pdf_bytes(text: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 750
    for line in text.split("\n"):
        if y < 60:
            c.showPage()
            y = 750
        c.drawString(40, y, line)
        y -= 14
    c.save()
    return buf.getvalue()


def _docx_bytes(text: str) -> bytes:
    import docx

    d = docx.Document()
    for line in text.split("\n"):
        if line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        else:
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_sections_detected():
    _, sections = split_sections(TEXT, [TEXT])
    assert [s.name for s in sections] == ["summary", "education", "skills", "projects", "experience",
                                          "certifications", "achievements"]


def test_personal_info_and_summary():
    r = parse_resume(TEXT)
    assert r.personal["name"] == "Arjun Mehta"
    assert r.personal["email"] == "arjun.mehta@example.com"
    assert "github.com/arjunmehta" in r.personal["links"]
    assert r.summary.startswith("Computer Science student")


def test_skills_with_categories_and_provenance():
    r = parse_resume(TEXT)
    by = {s["name"]: s for s in r.skills}
    assert by["Python"]["category"] == "programming_language"
    assert by["Qdrant"]["category"] == "database"
    assert by["AWS"]["category"] == "cloud"
    assert by["Natural Language Processing"]["source"]["section"] == "Skills"  # alias "NLP" canonicalised
    assert by["CLIP"]["source"]["section"] == "PROJECTS"  # found in a project, tagged with its section


def test_project_knowledge_model():
    r = parse_resume(TEXT)
    rag, bug = r.projects
    assert rag["name"] == "Multimodal RAG Platform"
    assert {"Python", "FastAPI", "Qdrant", "CLIP", "LangChain"} <= set(rag["technologies"])
    assert len(rag["responsibilities"]) == 3
    assert rag["challenges"] == []  # nothing invented
    assert rag["metrics"] == []
    assert bug["challenges"] == ["noisy, duplicated bug reports reduced accuracy"]
    assert bug["solutions"] == ["adding a deduplication step using cosine similarity before training"]
    assert any("87% accuracy" in m for m in bug["metrics"])
    assert bug["problem"].startswith("A tool to help developers triage")


def test_experience_certifications_education():
    r = parse_resume(TEXT)
    exp = r.experiences[0]
    assert exp["kind"] == "internship"
    assert exp["title"] == "Machine Learning Intern"
    assert exp["organization"] == "DataNest Analytics"
    assert (exp["start_date"], exp["end_date"]) == ("May 2025", "Jul 2025")
    assert [c["name"] for c in r.certifications] == ["AWS Certified Cloud Practitioner", "Deep Learning Specialization"]
    assert r.certifications[0]["issuer"] == "Amazon Web Services"
    edu = [i for i in r.items if i["kind"] == "education"][0]
    assert "Vellore Institute of Technology" in edu["title"] and edu["date"] == "2022 - 2026"


def test_no_headings_produces_warning():
    r = parse_resume("Jane Doe\njane@example.com\nI like building things with Python and Flask.")
    assert r.warnings and "No standard resume section" in r.warnings[0]


def test_pdf_extraction_and_parse():
    doc = extract_document(_pdf_bytes(TEXT), "resume.pdf")
    assert doc.kind == "pdf" and doc.page_count >= 1
    r = parse_resume(doc.text, doc.pages)
    assert {p["name"] for p in r.projects} == {"Multimodal RAG Platform", "AI Bug Analyzer"}


def test_docx_extraction_and_parse():
    doc = extract_document(_docx_bytes(TEXT), "resume.docx")
    assert doc.kind == "docx"
    r = parse_resume(doc.text, doc.pages)
    assert r.personal["name"] == "Arjun Mehta"
    assert len(r.projects) == 2


@pytest.mark.parametrize("data,code", [
    (b"", "empty"),
    (b"\x00\x01\x02binary", "unsupported_type"),
    (b"%PDF-1.4 /Encrypt something", "encrypted"),
    (b"%PDF-1.4 obj << /JavaScript (app.alert(1)) >>", "unsafe_content"),
    (b"PK\x03\x04garbage", "malformed"),
])
def test_rejects_unsafe_or_malformed(data, code):
    with pytest.raises(DocumentError) as e:
        extract_document(data, "x")
    assert e.value.code == code


def test_rejects_macro_docx():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("word/document.xml", "<x/>")
        z.writestr("word/vbaProject.bin", "macro")
    with pytest.raises(DocumentError) as e:
        extract_document(buf.getvalue(), "evil.docm")
    assert e.value.code == "unsafe_content"


def test_upload_size_limit():
    with pytest.raises(DocumentError) as e:
        extract_document(b"a" * 2000, "big.txt", max_bytes=1000)
    assert e.value.code == "too_large"


def test_chunking_carries_required_metadata():
    p = CandidateProfile(user_id="u1", name="A", summary="S", personal_verified=True)
    p.skills = [Skill(name="Python", category="programming_language", verified=True, position=0)]
    p.projects = [Project(id="p1", name="Proj", description="Does things.", technologies=["Python"],
                          responsibilities=["Built it"], challenges=["Hard part"], verified=False, position=0,
                          architecture=[], solutions=[], results=[], metrics=[], limitations=[], future_work=[],
                          testing=[], source={"document": "resume.pdf", "document_id": "d1", "page": 2})]
    p.experiences, p.certifications, p.items = [], [], []
    chunks = profile_chunks(p)
    fields = {c.meta.get("field") for c in chunks if c.collection == "projects"}
    assert fields == {"overview", "responsibilities", "challenges"}
    for c in chunks:
        for key in ("candidate_id", "document_type", "section", "verification_status", "source", "created_at"):
            assert key in c.meta, (key, c.meta)
    proj = next(c for c in chunks if c.collection == "projects")
    assert proj.meta["verification_status"] == "unverified" and proj.meta["page"] == 2
    assert proj.meta["project"] == "Proj" and proj.meta["technology"] == ["Python"]
    # deterministic point ids -> idempotent re-indexing
    assert proj.point_id("u1", 0) == proj.point_id("u1", 0)


def test_sentence_windows():
    text = " ".join(f"Sentence number {i} is here." for i in range(60))
    windows = sentence_windows(text, max_words=40)
    assert len(windows) > 3 and all(len(w.split()) <= 45 for w in windows)
