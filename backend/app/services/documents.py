"""Upload validation, safety checks and text extraction (PDF / DOCX / TXT).

Pipeline: validate -> safety check -> extract -> OCR fallback -> normalise.
Files are processed in memory; the only disk write is a private temporary
directory used by the optional OCR path, removed immediately afterwards.
"""

from __future__ import annotations

import hashlib
import io
import re
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field

from app.core.config import get_settings


class DocumentError(ValueError):
    """Raised for malformed, unsupported or unsafe uploads (-> HTTP 400/415)."""

    def __init__(self, message: str, code: str = "invalid_document") -> None:
        super().__init__(message)
        self.code = code


@dataclass
class ExtractedDocument:
    kind: str  # pdf | docx | txt
    content_type: str
    text: str
    pages: list[str]
    sha256: str
    size_bytes: int
    warnings: list[str] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)


_CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}
_MAX_ZIP_UNCOMPRESSED = 60 * 1024 * 1024
_MAX_ZIP_RATIO = 120
_PDF_ACTIVE_CONTENT = (b"/JavaScript", b"/JS ", b"/JS(", b"/Launch", b"/EmbeddedFile", b"/RichMedia")


def sniff_kind(data: bytes, filename: str) -> str:
    """Determine the real file type from content, not the client-supplied name."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "docx"
    if b"\x00" in data[:4096]:
        raise DocumentError("Binary file type is not supported. Upload PDF, DOCX or TXT.", "unsupported_type")
    return "txt"


def _check_pdf_safety(data: bytes) -> list[str]:
    warnings = []
    head = data[: 2 * 1024 * 1024]
    if b"/Encrypt" in head:
        raise DocumentError("Encrypted PDFs are not supported. Remove the password and re-upload.", "encrypted")
    for marker in _PDF_ACTIVE_CONTENT:
        if marker in data:
            raise DocumentError(
                "This PDF contains active content (scripts, launch actions or embedded files) and was rejected.",
                "unsafe_content",
            )
    return warnings


def _check_docx_safety(data: bytes) -> None:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError("The DOCX file is corrupt.", "malformed") from exc
    names = set(zf.namelist())
    if "word/document.xml" not in names or "[Content_Types].xml" not in names:
        raise DocumentError("ZIP archive is not a Word document.", "unsupported_type")
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise DocumentError("Macro-enabled documents are not accepted.", "unsafe_content")
    total = 0
    for info in zf.infolist():
        total += info.file_size
        if info.compress_size and info.file_size / max(info.compress_size, 1) > _MAX_ZIP_RATIO:
            raise DocumentError("Document compression ratio is suspicious (possible zip bomb).", "unsafe_content")
    if total > _MAX_ZIP_UNCOMPRESSED:
        raise DocumentError("Document is too large when decompressed.", "too_large")


_BULLETS = re.compile(r"^[\s]*[•●▪■‣⁃∙◦➢➔➜*–—]\s*", re.M)


def normalise_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = _BULLETS.sub("- ", text)
    text = re.sub(r"[  ]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _extract_pdf(data: bytes, warnings: list[str]) -> list[str]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentError("The PDF could not be read (malformed file).", "malformed") from exc
    if len(pages) > 30:
        raise DocumentError("PDF has too many pages for a resume or job description (max 30).", "too_large")
    if sum(len(p.strip()) for p in pages) < 40:
        ocr_pages = _ocr_pdf(data, warnings)
        if ocr_pages:
            return ocr_pages
    return pages


def _ocr_pdf(data: bytes, warnings: list[str]) -> list[str] | None:
    """OCR fallback for scanned PDFs. Requires the optional `ocr` extra plus
    the tesseract and poppler binaries; degrades to a clear warning otherwise."""
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except ImportError:
        warnings.append(
            "This PDF appears to be scanned (no text layer). OCR is not installed on this server; "
            "install the 'ocr' extra or upload a text-based PDF/DOCX."
        )
        return None
    try:
        with tempfile.TemporaryDirectory(prefix="iv-ocr-") as tmp:
            images = convert_from_bytes(data, dpi=250, output_folder=tmp, fmt="png")
            pages = [pytesseract.image_to_string(img) for img in images]
        warnings.append("Text was recovered with OCR; please review extracted details carefully.")
        return pages
    except Exception as exc:  # binaries missing / OCR failure
        warnings.append(f"OCR failed ({type(exc).__name__}); upload a text-based file instead.")
        return None


def _extract_docx(data: bytes) -> list[str]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("The DOCX file could not be read.", "malformed") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        if style.lower().startswith("list") and not text.startswith("-"):
            text = "- " + text
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return ["\n".join(lines)]


def _extract_txt(data: bytes) -> list[str]:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 always decodes
        raise DocumentError("Unable to decode text file.", "malformed")
    printable = sum(ch.isprintable() or ch in "\n\r\t" for ch in text[:5000])
    if text and printable / min(len(text), 5000) < 0.9:
        raise DocumentError("File does not look like plain text.", "unsupported_type")
    return [text]


def extract_document(data: bytes, filename: str, max_bytes: int | None = None) -> ExtractedDocument:
    max_bytes = max_bytes or get_settings().max_upload_mb * 1024 * 1024
    if not data:
        raise DocumentError("The uploaded file is empty.", "empty")
    if len(data) > max_bytes:
        raise DocumentError(f"File exceeds the {max_bytes // (1024 * 1024)} MB upload limit.", "too_large")
    if len(filename) > 255 or any(c in filename for c in "\x00/\\"):
        filename = re.sub(r"[\x00/\\]", "_", filename)[:255]

    kind = sniff_kind(data, filename)
    warnings: list[str] = []
    if kind == "pdf":
        warnings += _check_pdf_safety(data)
        pages = _extract_pdf(data, warnings)
    elif kind == "docx":
        _check_docx_safety(data)
        pages = _extract_docx(data)
    else:
        pages = _extract_txt(data)

    pages = [normalise_text(p) for p in pages]
    text = "\n\n".join(p for p in pages if p)
    if len(text.strip()) < 20:
        raise DocumentError(
            "No readable text could be extracted from this file." + (" " + warnings[-1] if warnings else ""),
            "no_text",
        )
    return ExtractedDocument(
        kind=kind,
        content_type=_CONTENT_TYPES[kind],
        text=text,
        pages=pages,
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        warnings=warnings,
    )


def page_of(pages: list[str], snippet: str) -> int:
    """1-based page on which `snippet` first appears (best effort)."""
    probe = snippet.strip()[:60]
    for i, page in enumerate(pages, start=1):
        if probe and probe in page:
            return i
    return 1
